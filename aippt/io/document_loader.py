# -*- coding: utf-8 -*-
"""从磁盘加载流水线输入：纯文本 / Markdown / Word docx。"""
# 从 pathlib 导入 Path
from pathlib import Path


def load_document_text(input_path: Path) -> str:
    """
    按扩展名加载文档正文（去除首尾空白）。

    @param input_path - 用户传入的输入文件路径
    @returns 合并后的 UTF-8 文本字符串
    @raises ValueError - 不支持的扩展名
    @raises ImportError - 缺少 python-docx 却传入 .docx
    """
    suffix = input_path.suffix.lower()  # 规范化后缀便于分支判断
    if suffix == ".docx":  # Microsoft Word Open XML 文稿
        return _extract_docx_plain_text(input_path)  # 走专用抽取逻辑
    if suffix in (".txt", ".md", ".markdown"):  # 纯文本或轻标记文档
        raw = input_path.read_text(encoding="utf-8-sig")  # utf-8-sig 吞掉 Windows 记事本 BOM
        return raw.strip()  # 去掉首尾空白与多余换行边缘情况
    msg = f"不支持的输入类型：{suffix}（当前支持 .txt .md .docx）：{input_path}"  # 构造可读异常说明
    raise ValueError(msg)  # 让 CLI 捕获并打印 stderr


def _extract_docx_plain_text(docx_path: Path) -> str:
    """
    使用 python-docx 抽取段落与表格单元格为可读纯文本。

    @param docx_path - .docx 文件路径
    @returns 以双换行分段的正文
    """
    try:
        # 延迟导入：未安装依赖时不要污染 txt-only 用户 import graph
        from docx import Document  # type: ignore[import-untyped] # noqa: WPS433
    except ImportError as err:
        msg = "读取 .docx 需要安装 python-docx：pip install python-docx"  # 安装指引
        raise ImportError(msg) from err  # 保留原始链路便于调试

    document = Document(str(docx_path))  # 打开 OPC 包
    chunks: list[str] = []  # 收集非空文本片段

    for paragraph in document.paragraphs:  # 顺序段落（保留大纲近似顺序）
        text = paragraph.text.strip()  # 去掉首尾空格
        if text:  # 跳过纯空段避免噪声
            chunks.append(text)  # 追加有效段落

    for table in document.tables:  # 遍历表格（许多公文正文写在表中）
        for row in table.rows:  # 逐行拼接单元格
            cells = [cell.text.strip() for cell in row.cells]  # 每格裁剪空白
            row_line = " | ".join(c for c in cells if c)  # 用竖线分隔格内文本
            if row_line:  # 非空行才写入
                chunks.append(row_line)  # 表格行作为一行追加

    merged = "\n\n".join(chunks).strip()  # 段落间留白增强模型可读性
    return merged  # 返回全文字符串
